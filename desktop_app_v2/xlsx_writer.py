"""Human-friendly 3-sheet xlsx, applying reviewer decisions."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline import LIST_FIELDS

HEADER = PatternFill("solid", fgColor="305496")
RED = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREEN = PatternFill("solid", fgColor="C6EFCE")
HFONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)

ALL_FIELDS = [
    "title", "subtitle", "creator", "contributor", "date_text", "date_year",
    "publisher", "scale_text", "scale_ratio", "projection", "edition",
    "coordinates_text", "bbox_west", "bbox_east", "bbox_south", "bbox_north",
    "country", "place_names", "province", "city", "district",
    "legend_content", "notes", "map_type", "subject", "coverage",
    "medium", "language", "condition", "identifier", "rights", "source",
    "relation", "has_insets", "description",
]


def _kept_list(field, items, filename, store):
    """Return display string for a list field, applying review decisions."""
    kept = []
    for it in items or []:
        if not isinstance(it, dict):
            continue
        if it.get("category", "").upper() != "INFERRED":
            kept.append(it["value"]); continue
        if store is None:
            continue
        dec = store.get_decision(f"{filename}|{field}|{it['value']}")
        if not dec:
            continue
        if dec["action"] == "remove":
            continue
        if dec["action"] == "edit" and dec.get("value"):
            kept.append(dec["value"])
        elif dec["action"] == "confirm":
            kept.append(it["value"])
    return ", ".join(kept)


def write_catalog(results, store, out_path: str):
    wb = Workbook()

    # Sheet 1: Review Queue
    ws = wb.active
    ws.title = "Review Queue"
    ws["A1"] = "🔍 Items to review"; ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A2:F2")
    ws["A2"] = ("Items where the AI couldn't find direct visual evidence. "
                "Confirm, correct, or remove each. Not in the main catalog until confirmed.")
    ws["A2"].alignment = WRAP
    heads = ["Map filename", "Field", "AI value", "AI reasoning", "Your decision", "Your notes"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(4, c, h); cell.fill = HEADER; cell.font = HFONT
    for c, w in enumerate([40, 16, 28, 50, 22, 36], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    dv = DataValidation(type="list",
        formula1='"Pending,Confirmed OK,Removed,Edited"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("E5:E10000")

    review_count = 0
    if store:
        r = 5
        for item in sorted(store.queue, key=lambda x: (x["filename"], x["field"])):
            review_count += 1
            dec = store.get_decision(item["key"])
            ws.cell(r, 1, item["filename"]).alignment = WRAP
            ws.cell(r, 2, item["field"])
            ws.cell(r, 3, item["ai_value"])
            ws.cell(r, 4, item["ai_evidence"]).alignment = WRAP
            if dec is None:
                txt, fill = "Pending", YELLOW
            elif dec["action"] == "remove":
                txt, fill = "Removed", RED
            elif dec["action"] == "edit":
                txt, fill = f"Edited → {dec.get('value','')}", GREEN
            else:
                txt, fill = "Confirmed OK", GREEN
            dc = ws.cell(r, 5, txt); dc.fill = fill
            ws.cell(r, 6, dec.get("note", "") if dec else "")
            r += 1
    ws.freeze_panes = "A5"

    # Sheet 2: Full Catalog
    ws2 = wb.create_sheet("Full Catalog")
    ws2["A1"] = "📊 Full extracted metadata (export this for Trove)"
    ws2["A1"].font = Font(bold=True, size=14)
    heads2 = ["Filename"] + ALL_FIELDS + ["AI status"]
    for c, h in enumerate(heads2, 1):
        cell = ws2.cell(3, c, h); cell.fill = HEADER; cell.font = HFONT
        ws2.column_dimensions[get_column_letter(c)].width = 22
    row = 4
    for res in results:
        ws2.cell(row, 1, res["filename"]).font = Font(bold=True)
        meta = res.get("metadata", {}) if res.get("ok") else {}
        for c, field in enumerate(ALL_FIELDS, 2):
            if field in LIST_FIELDS:
                val = _kept_list(field, meta.get(field), res["filename"], store)
            else:
                v = meta.get(field)
                val = "" if v is None else str(v)
            ws2.cell(row, c, val).alignment = WRAP
        st = ws2.cell(row, len(heads2),
                      "✓ OK" if res.get("ok") else f"✗ {str(res.get('error',''))[:60]}")
        if not res.get("ok"):
            st.fill = RED
        row += 1
    ws2.freeze_panes = "B4"

    # Sheet 3: AI Reasoning
    ws3 = wb.create_sheet("AI Reasoning")
    ws3["A1"] = "🔍 AI reasoning traces (optional)"; ws3["A1"].font = Font(bold=True, size=14)
    for c, h in enumerate(["Filename", "AI reasoning", "AI raw answer"], 1):
        cell = ws3.cell(3, c, h); cell.fill = HEADER; cell.font = HFONT
    ws3.column_dimensions["A"].width = 40
    ws3.column_dimensions["B"].width = 80
    ws3.column_dimensions["C"].width = 80
    r3 = 4
    for res in results:
        ws3.cell(r3, 1, res["filename"])
        ws3.cell(r3, 2, str(res.get("reasoning", ""))[:8000]).alignment = WRAP
        ws3.cell(r3, 3, str(res.get("raw_answer", ""))[:8000]).alignment = WRAP
        r3 += 1
    ws3.freeze_panes = "A4"

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"review_items": review_count, "total_maps": len(results)}
